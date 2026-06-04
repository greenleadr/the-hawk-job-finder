"""Export all jobs from jobs.db to an Excel (.xlsx) file.

Usage:
    python export_jobs.py [output_path]

Output defaults to jobs_export_YYYYMMDD.xlsx in the current directory.
"""

import json
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

_DB_PATH = Path(__file__).resolve().parent / "data" / "jobs.db"
_TRACKING_DB_PATH = Path(__file__).resolve().parent / "data" / "tracking.db"


def _fetch_jobs(jobs_db: Path, tracking_db: Path) -> list[dict]:
    conn = sqlite3.connect(str(jobs_db))
    conn.row_factory = sqlite3.Row

    # Attach tracking db if it exists
    if tracking_db.exists():
        conn.execute(f"ATTACH DATABASE '{tracking_db}' AS tracking")
        query = """
            SELECT
                j.id,
                j.title,
                j.company,
                j.location,
                j.url,
                j.source,
                j.date_posted,
                j.score,
                j.matched_skills,
                j.flags,
                j.status AS db_status,
                j.first_seen,
                j.last_seen,
                t.status AS tracked_status,
                t.notes,
                t.updated AS tracking_updated
            FROM jobs j
            LEFT JOIN tracking.tracking t ON t.job_id = j.id
            ORDER BY j.first_seen DESC
        """
    else:
        query = """
            SELECT
                id, title, company, location, url, source, date_posted,
                score, matched_skills, flags, status AS db_status,
                first_seen, last_seen,
                NULL AS tracked_status, NULL AS notes, NULL AS tracking_updated
            FROM jobs
            ORDER BY first_seen DESC
        """

    rows = conn.execute(query).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def _parse_json_list(value: str | None) -> str:
    if not value:
        return ""
    try:
        items = json.loads(value)
        if isinstance(items, list):
            return ", ".join(str(i) for i in items)
    except (json.JSONDecodeError, TypeError):
        pass
    return value or ""


COLUMNS = [
    ("ID",              "id",               18),
    ("Title",           "title",            40),
    ("Company",         "company",          25),
    ("Location",        "location",         25),
    ("URL",             "url",              50),
    ("Source",          "source",           14),
    ("Date Posted",     "date_posted",      16),
    ("Score",           "score",            8),
    ("Matched Skills",  "matched_skills",   40),
    ("Flags",           "flags",            35),
    ("DB Status",       "db_status",        14),
    ("First Seen",      "first_seen",       20),
    ("Last Seen",       "last_seen",        20),
    ("Tracked Status",  "tracked_status",   16),
    ("Notes",           "notes",            40),
    ("Tracking Updated","tracking_updated", 20),
]

HEADER_FILL = PatternFill("solid", fgColor="1A1A2E")
HEADER_FONT = Font(bold=True, color="E2B96A", size=11)
ALT_FILL    = PatternFill("solid", fgColor="F5F5F5")


def export(output_path: Path | None = None) -> Path:
    if output_path is None:
        stamp = datetime.now().strftime("%Y%m%d")
        output_path = Path(f"jobs_export_{stamp}.xlsx")

    jobs = _fetch_jobs(_DB_PATH, _TRACKING_DB_PATH)
    print(f"Exporting {len(jobs)} jobs → {output_path}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All Jobs"

    # Header row
    for col_idx, (header, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    # Data rows
    for row_idx, job in enumerate(jobs, start=2):
        fill = ALT_FILL if row_idx % 2 == 0 else None
        for col_idx, (_, field, _) in enumerate(COLUMNS, start=1):
            value = job.get(field)
            # Decode JSON arrays to comma-separated strings
            if field in ("matched_skills", "flags"):
                value = _parse_json_list(value)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if fill:
                cell.fill = fill
            # Make URL clickable
            if field == "url" and value:
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")

    # Auto-filter on header row
    ws.auto_filter.ref = ws.dimensions

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "Export generated"
    ws2["B1"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    ws2["A2"] = "Total jobs"
    ws2["B2"] = len(jobs)

    sources: dict[str, int] = {}
    statuses: dict[str, int] = {}
    for j in jobs:
        sources[j.get("source") or "unknown"] = sources.get(j.get("source") or "unknown", 0) + 1
        st = j.get("tracked_status") or j.get("db_status") or "open"
        statuses[st] = statuses.get(st, 0) + 1

    ws2["A4"] = "By Source"
    ws2["A4"].font = Font(bold=True)
    for i, (src, cnt) in enumerate(sorted(sources.items()), start=5):
        ws2.cell(row=i, column=1, value=src)
        ws2.cell(row=i, column=2, value=cnt)

    ws2["D4"] = "By Status"
    ws2["D4"].font = Font(bold=True)
    for i, (st, cnt) in enumerate(sorted(statuses.items()), start=5):
        ws2.cell(row=i, column=4, value=st)
        ws2.cell(row=i, column=5, value=cnt)

    for col in ["A", "B", "D", "E"]:
        ws2.column_dimensions[col].width = 20

    wb.save(output_path)
    print(f"Saved: {output_path.resolve()}")
    return output_path


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    export(out)
